#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""解析合并单元格式研究生课表 xlsx（教务总表）。

结构约定（教务系统打印版通用）：
  A列=时间段(上午1-4/下午5-8/晚上1 9-13/晚上2 14-15)  B列=节次锚点  C..I列=星期一~日
  单元格课程块用 ◇ 分隔信息：课程名 ◇ 周次(节次) ◇ 楼号 ◇ 教学班组成 ◇ 任课教师
  一个格子里多门课 = 不同教学班；块物理行号≠真实节次，跨度以块内 (x-y节) 为准

用法:
  python3 parse_timetable_xlsx.py <课表.xlsx>
  python3 parse_timetable_xlsx.py <课表.xlsx> 专硕        # 按关键词过滤教学班/课名

需要: pip3 install openpyxl
"""
import re
import sys
import openpyxl


def parse_cell(text):
    """把一个单元格文本拆成课程块列表"""
    if not text:
        return []
    text = text.strip()
    # 换行后紧跟非◇开头的行 = 新课程块起点（也可能是课名被单元格换行拆开，需人工复核）
    parts = re.split(r'\n(?=[^\n◇])', text)
    blocks = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        m = re.match(r'^(.+?)\n◇(.+?)◇\s*(.*?)\s*$', p, re.S)
        if not m:
            m = re.match(r'^(.+?)\n◇(.+?)◇', p, re.S)
            if not m:
                blocks.append({'name': p[:50], 'weeks': '?', 'audience': '', 'teacher': ''})
                continue
            name, weeks, rest = m.group(1).strip(), m.group(2).strip(), p[m.end():].strip()
        else:
            name, weeks, rest = m.group(1).strip(), m.group(2).strip(), m.group(3).strip()
        audience, teacher = '', ''
        for line in rest.split('\n'):
            line = line.strip()
            if line.startswith('教学班组成'):
                audience = line.split('：', 1)[-1]
            elif line.startswith('任课教师'):
                teacher = line.split('：', 1)[-1]
        blocks.append({'name': name, 'weeks': weeks, 'audience': audience, 'teacher': teacher})
    return blocks


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    path = sys.argv[1]
    keyword = sys.argv[2] if len(sys.argv) > 2 else None
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    print(f'# {ws.title}  rows={ws.max_row} cols={ws.max_column}')
    for d_idx in range(3, 10):
        day = ws.cell(row=2, column=d_idx).value
        if not day:
            continue
        print(f'\n== {day} ==')
        found = False
        for r in range(3, ws.max_row + 1):
            for b in parse_cell(ws.cell(row=r, column=d_idx).value):
                text = '  [%s] %s | %s | %s' % (b['weeks'], b['name'], b['audience'], b['teacher'])
                if keyword and keyword not in (b['name'] + b['audience']):
                    continue
                print(text)
                found = True
        if not found:
            print('  (无课)')


if __name__ == '__main__':
    main()
